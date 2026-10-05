"""Unit tests: run with  python -m pytest -q"""
import asyncio
import json
import os
import re
import sys
import urllib.parse
from types import SimpleNamespace

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("ANTHROPIC_API_KEY", "test")
import teams_meeting_bot as tb  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def wake(text):
    m = tb.WAKE.search(text) or tb.WAKE_START.search(text)
    return text[m.end():].strip() if m else None


def test_wake_phrase_variants():
    assert wake("Hey Claude, what did Mark commit to?") == "what did Mark commit to?"
    assert wake("Okay cloud what's the timeline") == "what's the timeline"   # caption mishear
    assert wake("Claude, summarize that.") == "summarize that."
    assert wake("Hey, Claude.") == ""


def test_wake_phrase_negatives():
    assert wake("We moved everything to the Azure cloud last year.") is None
    assert wake("hey team, cloud costs are up") is None
    assert wake("I asked Claude yesterday about it") is None


def test_teams_web_url_forces_browser_join():
    url = "https://teams.microsoft.com/l/meetup-join/19%3ameeting_abc%40thread.v2/0?context=%7b%7d"
    q = dict(urllib.parse.parse_qsl(urllib.parse.urlparse(tb.teams_web_url(url)).query))
    assert q["msLaunch"] == "false" and q["suppressPrompt"] == "true" and q["context"] == "{}"


def make_bot(tmp_path, *extra):
    args = tb.parse_args(["--url", "x", "--outdir", str(tmp_path), *extra])
    return tb.MeetingBot(args)


def test_finalizer_emits_new_words_only(tmp_path):
    bot = make_bot(tmp_path, "--no-agent", "--stable", "0.05")
    heard = []

    async def fake_wake(speaker, text):
        heard.append((speaker, text))
    bot.check_wake = fake_wake

    async def run():
        task = asyncio.create_task(bot.finalizer())
        await bot.on_caption(None, {"id": 1, "name": "Mark Chen", "text": "I'll own the renegotiation."})
        await asyncio.sleep(0.4)
        await bot.on_caption(None, {"id": 1, "name": "Mark Chen",
                                    "text": "I'll own the renegotiation. And loop in finance."})
        await asyncio.sleep(0.4)
        bot.ended.set()
        await task
    asyncio.run(run())
    assert heard == [("Mark Chen", "I'll own the renegotiation."),
                     ("Mark Chen", "And loop in finance.")]


def test_agent_options_are_locked_down():
    ag = tb.MeetingAgent(os.path.join(ROOT, "agent"), "claude-sonnet-5-5", 12, 2.0, 120)
    o = ag.options()
    assert o.tools == ["Read", "Grep", "Glob"]                 # no Bash/Write/Edit/WebFetch
    assert o.permission_mode == "dontAsk"
    assert o.allowed_tools == []          # reads in cwd need no rule; a bare "Read" allows any file
    assert ag.servers == [] and o.mcp_servers == {}            # no MCP server ships enabled
    assert o.cwd.endswith(os.path.join("agent", "project"))
    assert o.setting_sources == [] and o.max_budget_usd == 2.0
    assert o.verbatim_prompts                                  # no @file expansion in captions
    assert o.extra_args == {"no-session-persistence": None}    # no transcript copies on disk
    assert o.cli_path == tb.bundled_engine() and o.cli_path    # the pinned engine, not PATH's
    assert o.stderr is tb.engine_note                          # engine notices: log details
    assert "AI meeting assistant for the leadership team" in o.system_prompt


def test_engine_notes_are_technical_details(capsys):
    tb.engine_note('[claude-code:unrecognized_model] {"model":"my-deployment"}\n')
    tb.engine_note("   \n")
    out = capsys.readouterr().out.splitlines()
    assert len(out) == 1 and re.match(
        r'\[\d\d:\d\d:\d\d\]   engine: \[claude-code:unrecognized_model\] \{"model":"my-deployment"\}$', out[0])


def test_mcp_servers_are_the_only_allow_rules(tmp_path):
    (tmp_path / "project").mkdir()
    (tmp_path / "instructions.md").write_text("x", encoding="utf-8")
    (tmp_path / "mcp.json").write_text('{"mcpServers": {"docs": {"type": "http", "url": "https://x"}}}')
    o = tb.MeetingAgent(str(tmp_path), "m", 12, 2.0, 120).options()
    assert o.allowed_tools == ["mcp__docs"] and str(o.mcp_servers).endswith("mcp.json")
    assert tb.MeetingAgent(str(tmp_path), "m", 12, 2.0, 120).options(use_tools=False).allowed_tools == []


def test_example_mcp_config_is_valid_json():
    with open(os.path.join(ROOT, "agent", "mcp.example.json"), encoding="utf-8") as f:
        assert "microsoft-learn" in json.load(f)["mcpServers"]


def test_agent_prompt_contains_transcript_and_question():
    p = tb.AGENT_PROMPT.format(transcript="[10:00] Dana: hi", asker="Dana", question="what's capex?",
                               now=tb.now_text())
    assert "[10:00] Dana: hi" in p and "Dana just asked you" in p and "what's capex?" in p
    assert "Current date and time: " + tb.dt.datetime.now().strftime("%A") in p


def test_chat_text_strips_markdown():
    assert tb.chat_text("## Answer\n- **Single meeting:** use `RSC`") == "Answer\n- Single meeting: use RSC"


def test_office_materials_become_searchable_text(tmp_path):
    import docx, pptx, openpyxl
    d = docx.Document(); d.add_paragraph("Vendor shortlist: Fabrikam and Northwind."); d.save(tmp_path / "brief.docx")
    pr = pptx.Presentation(); s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = "Roadmap"; s.placeholders[1].text = "Phase 2 starts January 12"; pr.save(tmp_path / "deck.pptx")
    wb = openpyxl.Workbook(); wb.active.append(["Line", "Cost"]); wb.active.append(["UPS", 420000]); wb.save(tmp_path / "budget.xlsx")
    assert tb.prepare_materials(str(tmp_path)) == 3
    out = tmp_path / tb.TEXT_DIR
    assert "Fabrikam" in (out / "brief.docx.md").read_text()
    assert "Phase 2 starts January 12" in (out / "deck.pptx.md").read_text()
    assert "UPS | 420000" in (out / "budget.xlsx.md").read_text()
    assert tb.prepare_materials(str(tmp_path)) == 0          # unchanged files are not redone


def test_pdf_materials_become_searchable_text(tmp_path, make_pdf, capsys):
    make_pdf(tmp_path / "budget.pdf", "Plant 3 network refresh: 780K, owner Priya Nair")
    make_pdf(tmp_path / "scan.pdf", None)                    # image only, like a scanned page
    assert tb.prepare_materials(str(tmp_path)) == 1          # only budget.pdf has text to search
    out = tmp_path / tb.TEXT_DIR
    assert "## Page 1\nPlant 3 network refresh: 780K, owner Priya Nair" in (out / "budget.pdf.md").read_text()
    assert tb.SCANNED_NOTE in (out / "scan.pdf.md").read_text()
    log = capsys.readouterr().out
    assert "Preparing the knowledge folder (2 new or changed file(s))..." in log
    assert "could not find text in scan.pdf" in log
    assert "Knowledge folder ready (1 of 2 file(s) converted to searchable text)." in log
    assert tb.prepare_materials(str(tmp_path)) == 0          # nothing new: no work, no messages
    assert capsys.readouterr().out == ""


def test_materials_follow_edits_replacements_and_deletions(tmp_path, make_pdf, capsys, monkeypatch):
    pdf, copy = tmp_path / "brief.pdf", tmp_path / tb.TEXT_DIR / "brief.pdf.md"

    def write(text, age_s):
        make_pdf(pdf, text)
        t = pdf.stat().st_mtime + age_s
        os.utime(pdf, (t, t))

    write("Version one", 0)
    assert tb.prepare_materials(str(tmp_path)) == 1
    write("Version two", 60)                                 # edited later
    assert tb.prepare_materials(str(tmp_path)) == 1 and "Version two" in copy.read_text()
    write("Version three", -86400)                           # an older copy dropped over it
    assert tb.prepare_materials(str(tmp_path)) == 1 and "Version three" in copy.read_text()
    pdf.write_bytes(b"not a pdf")                            # unreadable: the outdated text goes
    assert tb.prepare_materials(str(tmp_path)) == 0
    assert "Version three" not in copy.read_text() and "Could not extract text" in copy.read_text()
    assert "could not convert brief.pdf" in capsys.readouterr().out
    assert tb.prepare_materials(str(tmp_path)) == 0          # not retried until the file changes
    assert capsys.readouterr().out == ""

    sub = tmp_path / "prior meetings"
    sub.mkdir()
    make_pdf(sub / "agenda.pdf", "Item 1: Q4 budget")

    def no_pypdf(path):
        raise ModuleNotFoundError("No module named 'pypdf'", name="pypdf")
    monkeypatch.setattr(tb, "_to_text", no_pypdf)            # setup.ps1 not re-run on this PC
    assert tb.prepare_materials(str(tmp_path)) == 0
    monkeypatch.undo()
    assert "could not convert agenda.pdf: pypdf is not installed (run setup.ps1)" in capsys.readouterr().out
    assert tb.prepare_materials(str(tmp_path)) == 1          # retried, not remembered as unreadable
    (sub / "agenda.pdf").unlink()
    make_pdf(sub / "minutes.pdf", "Decision: approve the UPS replacement")
    assert tb.prepare_materials(str(tmp_path)) == 1
    (sub / "minutes.pdf").unlink()                           # deleted: nothing of it stays searchable
    tb.prepare_materials(str(tmp_path))
    assert not (tmp_path / tb.TEXT_DIR / "prior meetings").exists()


def test_project_override_and_foundry_pinning(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAUDE_CODE_USE_FOUNDRY", "1")
    ag = tb.MeetingAgent(os.path.join(ROOT, "agent"), "my-opus-deployment", 12, 0.75, 120, str(tmp_path))
    o = ag.options()
    assert o.cwd == str(tmp_path)
    assert o.env["ANTHROPIC_DEFAULT_HAIKU_MODEL"] == "my-opus-deployment"
    assert o.env["CLAUDE_CODE_DISABLE_NONESSENTIAL_TRAFFIC"] == "1"
    assert o.env["ANTHROPIC_DEFAULT_SONNET_MODEL"] == "my-opus-deployment"
    assert "Microsoft Foundry" in ag.describe()
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_RESOURCE", "contoso-foundry")
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_API_KEY", "x")
    assert type(tb.make_client()).__name__ == "AsyncAnthropicFoundry"


def test_direct_join_url_skips_launcher():
    launcher = ("https://teams.microsoft.com/dl/launcher/launcher.html?url=%2F_%23%2Fmeet%2F210136600936824"
                "%3Fp%3DfGnTBFTncBYv6FJfHx%26anon%3Dtrue&type=meet&directDl=true&msLaunch=true")
    assert tb.direct_join_url(launcher) == \
        "https://teams.microsoft.com/_#/meet/210136600936824?p=fGnTBFTncBYv6FJfHx&anon=true"
    no_anon = "https://teams.microsoft.com/dl/launcher/launcher.html?url=%2F_%23%2Fmeet%2F123%3Fp%3Dabc"
    assert tb.direct_join_url(no_anon).endswith("/_#/meet/123?p=abc&anon=true")
    assert tb.direct_join_url("https://teams.microsoft.com/v2/") is None


def test_clean_name_removes_characters_teams_rejects():
    assert tb.clean_name("Claude (Meeting Assistant)") == "Claude Meeting Assistant"
    assert tb.clean_name("Ann's A.I. - Notes_Bot@HQ") == "Ann's A.I. - Notes_Bot@HQ"
    assert tb.clean_name("((()))") == "Claude Meeting Assistant"
    assert tb.base_name("Claude Meeting Assistant (Guest)") == "claude meeting assistant"


def test_direct_web_join_never_touches_launcher():
    meet = "https://teams.microsoft.com/meet/210136600936824?p=fGnTBFTncBYv6FJfHx"
    assert tb.direct_web_join(meet) == \
        "https://teams.microsoft.com/_#/meet/210136600936824?p=fGnTBFTncBYv6FJfHx&anon=true"
    assert tb.direct_web_join(meet, guest=False) == \
        "https://teams.microsoft.com/_#/meet/210136600936824?p=fGnTBFTncBYv6FJfHx"
    legacy = ("https://teams.microsoft.com/l/meetup-join/19%3ameeting_abc%40thread.v2/0"
              "?context=%7b%22Tid%22%3a%22t1%22%7d")
    d = tb.direct_web_join(legacy)
    assert d.startswith("https://teams.microsoft.com/_#/l/meetup-join/19%3ameeting_abc%40thread.v2/0?")
    assert d.endswith("&anon=true") and '{"Tid":"t1"}' in d
    assert tb.direct_web_join("https://example.com/meet/1") is None
    assert tb.direct_web_join("file:///x/mock.html") is None


def test_spoken_version_is_short_and_clean():
    ans = ("The Q4 IT capex budget is $2.4M. Priya Nair owns the Plant 3 refresh. "
           "The brief lists it at $780K, with a 10-week switch lead time and other detail "
           "that would take far too long to read aloud in a meeting setting at all. "
           "(source: q4_it_budget_brief.md) See https://x.y/z")
    sp = tb.spoken_version(ans)
    assert "source" not in sp and "http" not in sp
    assert sp == "The Q4 IT capex budget is $2.4M. Priya Nair owns the Plant 3 refresh."
    assert len(sp.split()) <= 25


def test_voice_modes(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    assert bot.voice_mode == "asked"
    assert bot.wants_voice("tell us out loud what the budget is")
    assert not bot.wants_voice("what is the budget")
    bot.voice_mode = "always"; assert bot.wants_voice("what is the budget")
    bot.voice_mode = "off"; assert not bot.wants_voice("say it out loud")
    assert bot.a.name == "Claude Meeting Assistant"


def test_azure_tts_request(monkeypatch):
    seen = {}

    class R:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def read(self): return b"RIFFwav"

    def fake_urlopen(req, timeout=0):
        seen.update(url=req.full_url, headers=dict(req.headers), body=req.data.decode())
        return R()
    monkeypatch.setattr(tb.urllib.request, "urlopen", fake_urlopen)
    assert tb.azure_tts("A & B <ok>", "en-US-AndrewNeural", "res1", "k1") == b"RIFFwav"
    assert seen["url"] == "https://res1.cognitiveservices.azure.com/tts/cognitiveservices/v1"
    assert seen["headers"]["Ocp-apim-subscription-key"] == "k1"
    assert "A &amp; B &lt;ok&gt;" in seen["body"] and "en-US-AndrewNeural" in seen["body"]


def test_pause_mid_question_is_one_question(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    asked = []
    bot.dispatch = lambda p: (asked.append(p["text"]), setattr(bot, "pending", None))

    async def run():
        await bot.check_wake("Scott Hay", "Hey Claude, what's the capital")
        t = bot.pending["t"]
        assert bot.due_question(t + 1.0) is None                 # still inside the pause
        await bot.check_wake("Scott Hay", "of South Africa out loud?")
        await bot.check_wake("Scott Hay", "Hey Claude what's the capital of South Africa out loud?")
        assert bot.pending["text"] == "what's the capital of South Africa out loud?"   # re-render, no dup
        bot.caps["x"] = {"name": "Scott Hay", "dirty": True}
        assert bot.due_question(bot.pending["t"] + 5) is None    # asker still talking
        bot.caps["x"]["dirty"] = False
        p = bot.due_question(bot.pending["t"] + 0.7)             # ends with '?': short wait
        bot.dispatch(p)
    asyncio.run(run())
    assert asked == ["what's the capital of South Africa out loud?"]


def test_question_without_question_mark_waits_longer(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")

    async def run():
        await bot.check_wake("Dana", "Hey Claude summarize the budget discussion")
        t = bot.pending["t"]
        assert bot.due_question(t + 1.0) is None
        assert bot.due_question(t + 1.9)["text"] == "summarize the budget discussion"
        await bot.check_wake("Mark", "unrelated words from someone else")
        assert bot.pending["speaker"] == "Dana"
    asyncio.run(run())


def test_duplicate_question_is_answered_once(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    spawned = []
    bot.spawn = lambda coro: (spawned.append(1), coro.close())
    now = tb.time.monotonic()
    bot.dispatch({"speaker": "Dana", "text": "what's the capital", "t": now, "started": now})
    bot.dispatch({"speaker": "Dana", "text": "What's the capital of South Africa?", "t": now, "started": now})
    bot.dispatch({"speaker": "Mark", "text": "what's the capital", "t": now, "started": now})
    assert len(spawned) == 2


def test_repeat_requests_are_recognised():
    yes = ["can you say that out loud?", "Say that out loud.", "read it out loud please",
           "repeat that", "could you read your answer aloud", "say it again", "okay read that back",
           "Can you speak that?", "read the answer out loud"]
    no = ["say what the Q4 budget is out loud", "read the agenda out loud",
          "repeat what Mark said about the budget", "how many days are left in Q4?",
          "tell us out loud who owns Plant 3"]
    for q in yes:
        assert tb.REPEAT_ASK.match(q), q
    for q in no:
        assert not tb.REPEAT_ASK.match(q), q


def test_repeat_speaks_last_answer_without_agent(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    spoken, chat = [], []

    async def speak(t): spoken.append(t)
    async def post(t): chat.append(t)
    bot.speak, bot.post_chat = speak, post
    bot.last_answer = "There are 89 days left in Q4. (source: calendar)"
    asyncio.run(bot.answer("can you say that out loud?", "Scott Hay", to_chat=True))
    assert spoken == ["There are 89 days left in Q4."] and chat == []
    bot.voice_mode = "off"
    asyncio.run(bot.answer("say that out loud", "Scott Hay", to_chat=True))
    assert "turned off" in chat[0] and len(spoken) == 1


# "Ask privately" (app) / plain console text: the operator's question and its answer.
PRIVATE_Q, PRIVATE_A = "What is our walk-away price for Northwind?", "Walk away above 1.2 million."
REPLIES = [PRIVATE_A,                                       # private answer
           "Legal is reviewing the Northwind contract.",    # public answer
           "## Summary\nNorthwind is in legal review.",     # summary file
           "Northwind: legal review continues."]            # chat wrap-up


def fake_client(bot, sent, replies):
    """The direct-API model: records everything it is sent."""
    async def create(**kw):
        sent.append("\n".join([kw["system"]] + [c["text"] for c in kw["messages"][0]["content"]]))
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=replies.pop(0))])
    bot.client = SimpleNamespace(messages=SimpleNamespace(create=create))


def meeting_with_a_private_question(bot):
    """Dana speaks, the operator asks privately, Dana asks in the meeting, the meeting wraps
    up. Returns the meeting chat."""
    chat = []

    async def post(t): chat.append(t)
    bot.post_chat = post

    async def meeting():
        bot.add_line("Dana Ruiz", "Next item is the Northwind contract.")
        await bot.answer(PRIVATE_Q, "Operator", to_chat=False)
        await bot.answer("where are we on the Northwind contract?", "Dana Ruiz", to_chat=True)
        await bot.summarize(to_chat=True)
    asyncio.run(meeting())
    return chat


def assert_attendees_never_see_it(sent, chat):
    private_request, *later = sent
    assert "walk-away" in private_request              # the private question did go to the model
    assert len(later) == 3 and len(chat) == 2          # public answer, summary, wrap-up; 2 chat posts
    leaked = [name for name, request in zip(("public answer", "summary", "chat wrap-up"), later)
              if "walk-away" in request or "1.2 million" in request]
    assert leaked == []


def test_private_question_never_reaches_the_meeting(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    sent = []
    fake_client(bot, sent, list(REPLIES))
    assert_attendees_never_see_it(sent, meeting_with_a_private_question(bot))


def test_private_question_never_reaches_the_meeting_through_the_agent(tmp_path):
    bot = make_bot(tmp_path)                           # the agent path real meetings use
    sent, replies = [], list(REPLIES)

    async def ask(transcript, asker, question):
        sent.append(f"{transcript}\n{asker} asks: {question}")
        return replies.pop(0), [], 0.0

    async def complete(transcript, task, timeout=None, budget=None):
        sent.append(f"{transcript}\n{task}")
        return replies.pop(0)
    bot.agent.ask, bot.agent.complete = ask, complete
    assert_attendees_never_see_it(sent, meeting_with_a_private_question(bot))


def test_private_follow_up_sees_earlier_private_answers(tmp_path):
    bot = make_bot(tmp_path)
    seen = []

    async def ask(transcript, asker, question):
        seen.append(transcript)
        return PRIVATE_A, [], 0.0
    bot.agent.ask = ask

    async def meeting():
        bot.add_line("Dana Ruiz", "Next item is the Northwind contract.")
        await bot.answer(PRIVATE_Q, "Operator", to_chat=False)
        await bot.answer("and who signs off on that?", "Operator", to_chat=False)
    asyncio.run(meeting())
    assert "Next item is the Northwind contract." in seen[1]       # the meeting so far
    assert "walk-away" in seen[1] and "1.2 million" in seen[1]     # plus the earlier private Q&A


def test_private_answers_are_saved_to_the_operator_file(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    fake_client(bot, [], [PRIVATE_A])
    bot.add_line("Dana Ruiz", "Next item is the Northwind contract.")
    asyncio.run(bot.answer(PRIVATE_Q, "Operator", to_chat=False))
    meeting_record = next(tmp_path.glob("meeting_transcript_*.md")).read_text(encoding="utf-8")
    operator_files = [p.read_text(encoding="utf-8") for p in tmp_path.glob("operator_private_*.md")]
    assert "Northwind contract." in meeting_record and "walk-away" not in meeting_record
    assert len(operator_files) == 1 and PRIVATE_Q in operator_files[0] and PRIVATE_A in operator_files[0]


def test_private_questions_alone_get_no_summary_or_wrap_up(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    sent, chat = [], []
    fake_client(bot, sent, [PRIVATE_A, "summary", "wrap-up"])

    async def post(t): chat.append(t)
    bot.post_chat = post
    asyncio.run(bot.answer(PRIVATE_Q, "Operator", to_chat=False))
    asyncio.run(bot.summarize(to_chat=True))
    assert len(sent) == 1 and chat == []               # only the private answer used the model


# A question heard in the meeting gets an immediate "Got it" that repeats it as heard: the asker
# knows Claude is on it, and a caption mishear shows before the answer does.
QUESTION, ANSWER = "what's the Q4 capex number?", "Q4 IT capex is $2.4M."
GOT_IT = f'Claude: Got it, Dana Ruiz – "{QUESTION}" Working on it.'


def record_chat(bot):
    """The meeting chat, in posting order."""
    chat = []

    async def post(t): chat.append(t)
    bot.post_chat = post
    return chat


def answer_instantly(bot):
    if bot.agent:
        async def ask(transcript, asker, q): return ANSWER, [], 0.0     # never yields to the loop
        bot.agent.ask = ask
    else:
        fake_client(bot, [], [ANSWER])


async def finish_spawned():
    """Let everything the bot spawned (answers, acknowledgements) run to the end."""
    while tasks := [t for t in asyncio.all_tasks() if t is not asyncio.current_task() and not t.done()]:
        await asyncio.wait(tasks)


def hear(bot, question=QUESTION, speaker="Dana Ruiz"):
    """The bot hears a question in the meeting and handles it to the end."""
    async def meeting():
        now = tb.time.monotonic()
        bot.dispatch({"speaker": speaker, "text": question, "t": now, "started": now})
        await finish_spawned()
    asyncio.run(meeting())


@pytest.mark.parametrize("extra", [(), ("--no-agent",)], ids=["agent", "no-agent"])
def test_heard_question_is_acknowledged_before_the_answer(tmp_path, extra):
    bot = make_bot(tmp_path, *extra)
    chat = record_chat(bot)
    answer_instantly(bot)
    hear(bot)
    assert chat == [GOT_IT, "Claude: " + ANSWER]


def test_acknowledgement_is_in_chat_while_the_agent_works(tmp_path):
    bot = make_bot(tmp_path)
    chat, seen_while_working = record_chat(bot), []

    async def ask(transcript, asker, q):
        await asyncio.sleep(0.05)
        seen_while_working.extend(chat)
        return ANSWER, [], 0.0
    bot.agent.ask = ask
    hear(bot)
    assert seen_while_working == [GOT_IT]


def test_long_question_is_shortened_in_the_acknowledgement(tmp_path):
    bot = make_bot(tmp_path)
    chat = record_chat(bot)
    answer_instantly(bot)
    q = "walk us through " + "the vendor consolidation plan and the network refresh " * 5
    hear(bot, q, "Mark Chen")
    assert chat[0].startswith('Claude: Got it, Mark Chen – "')
    heard = chat[0].split('"')[1]
    assert len(heard) <= 120 and heard.endswith("...") and q.startswith(heard[:-3])


def test_asker_is_told_when_no_answer_comes(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    chat = record_chat(bot)

    async def create(**kw): raise RuntimeError("model unreachable")
    bot.client = SimpleNamespace(messages=SimpleNamespace(create=create))
    hear(bot)
    assert chat == [GOT_IT, "Claude: Sorry, Dana Ruiz – I couldn't get an answer just now."]


@pytest.mark.parametrize("limit", ["Reached maximum budget ($2)", "Reached maximum number of turns (12)",
                                   "agent timed out"])
def test_a_tripped_limit_still_gets_an_answer(tmp_path, limit):
    """Per-question limits only stop the agent's file search: the asker still gets an answer."""
    bot = make_bot(tmp_path)
    chat = record_chat(bot)

    async def ask(transcript, asker, q): raise RuntimeError(limit)

    async def complete(transcript, task, timeout=None, budget=None): return ANSWER
    bot.agent.ask, bot.agent.complete = ask, complete
    hear(bot)
    assert chat == [GOT_IT, "Claude: " + ANSWER]


def test_summary_and_wrap_up_get_their_own_spend_limit(tmp_path):
    bot = make_bot(tmp_path)
    record_chat(bot)
    budgets = []

    async def complete(transcript, task, timeout=None, budget=None):
        budgets.append(budget)
        return "## Summary\nDone."
    bot.agent.complete = complete
    bot.add_line("Dana Ruiz", "We approved the network refresh.")
    asyncio.run(bot.summarize(to_chat=True))
    assert budgets == [tb.SUMMARY_BUDGET_USD] * 2 and tb.SUMMARY_BUDGET_USD > 2.0


def test_acknowledgement_can_be_turned_off(tmp_path):
    bot = make_bot(tmp_path, "--no-ack")
    chat = record_chat(bot)
    answer_instantly(bot)
    hear(bot)
    assert chat == ["Claude: " + ANSWER]


def test_operator_chat_question_gets_no_acknowledgement(tmp_path):
    bot = make_bot(tmp_path, "--no-agent")
    chat = record_chat(bot)
    answer_instantly(bot)
    bot.listening = True

    async def run():
        await bot.handle_command(f"/chat {QUESTION}")
        await finish_spawned()
    asyncio.run(run())
    assert chat == ["Claude: " + ANSWER]
